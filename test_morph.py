from pptx import Presentation
from pptx.oxml import parse_xml

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

morph_xml = """<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="med"><p14:morph option="byObject"/></p:transition>"""

transition_elem = parse_xml(morph_xml)
slide.element.insert(-1, transition_elem)
prs.save(r"d:\Final Project\fix\test_morph.pptx")
print("Morph XML parsed and inserted successfully!")
